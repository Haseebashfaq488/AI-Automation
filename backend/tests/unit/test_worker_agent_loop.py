import asyncio
import json

import pytest
from docx import Document

from app.registry import registry
from app.workers.base.contract import TaskContract
from app.workers.base.session import WorkerSession
from app.workers.base.engine import WorkerEngine


@pytest.fixture(autouse=True)
def sessions_root(tmp_path, monkeypatch):
    monkeypatch.setattr(WorkerSession, "SESSIONS_ROOT", tmp_path / "sessions")
    return tmp_path / "sessions"


@pytest.fixture
def docx_path(tmp_path):
    doc = Document()
    doc.add_paragraph("# Title")
    doc.add_paragraph("Body.")
    path = tmp_path / "report.docx"
    doc.save(str(path))
    return path


def make_contract(tmp_path, tools, max_steps=5):
    return TaskContract(
        objective="Fix the report headings",
        requirements=[],
        constraints=[],
        success_criteria=[],
        fs_scope=str(tmp_path),
        allowed_tools=tools,
        max_steps=max_steps,
    )


class ScriptedAgent:
    """Fake worker agent that replays scripted decisions and records steps."""

    def __init__(self, decisions):
        self.decisions = list(decisions)
        self.available = True
        self.recorded = []

    async def decide_next_step(self):
        if not self.decisions:
            return {"action": "done", "summary": "scripted finish"}
        return self.decisions.pop(0)

    def record_step(self, tool, success, output=None, error=None):
        self.recorded.append((tool, success))

    def inject_intervention(self, message):
        self.recorded.append(("INTERVENTION", message))


async def wait_done(engine, timeout=10):
    async def poll():
        while True:
            state = engine.get_state()
            if state["status"] in ("completed", "cancelled"):
                return state
            await asyncio.sleep(0.05)

    return await asyncio.wait_for(poll(), timeout)


# ── agent loop happy path ───────────────────────────────────────────────


@pytest.mark.asyncio
async def test_engine_agent_loop_runs_tools_and_finishes(sessions_root, tmp_path, docx_path):
    agent = ScriptedAgent(
        [
            {"action": "tool", "tool": "inspect_docx", "params": {"path": str(docx_path)}},
            {"action": "tool", "tool": "normalize_headings", "params": {"path": str(docx_path)}},
            {"action": "done", "summary": "Headings normalized."},
        ]
    )
    eng = WorkerEngine(registry, agent_factory=lambda contract: agent)
    await eng.run(make_contract(tmp_path, ["inspect_docx", "normalize_headings"]))
    await wait_done(eng)

    result = eng.get_result()
    assert result["success"] is True
    assert result["summary"] == "Headings normalized."
    assert result["tools_used"] == ["inspect_docx", "normalize_headings"]
    assert result["cancelled"] is False

    # file actually modified by the loop
    from docx import Document as D
    assert D(str(docx_path)).paragraphs[0].style.name == "Heading 1"

    # events.jsonl captured the run
    lines = (eng._session.path / "events.jsonl").read_text(encoding="utf-8").splitlines()
    types = [json.loads(l)["type"] for l in lines]
    assert types[0] == "WORK_STARTED"
    assert types[-1] == "WORK_COMPLETED"
    assert "STEP_COMPLETED" in types


# ── failures ────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_engine_agent_loop_fails_on_bad_decision(sessions_root, tmp_path):
    agent = ScriptedAgent([{"action": "fail", "reason": "LLM output unparseable"}])
    eng = WorkerEngine(registry, agent_factory=lambda contract: agent)
    await eng.run(make_contract(tmp_path, ["inspect_docx"]))
    await wait_done(eng)

    result = eng.get_result()
    assert result["success"] is False
    assert any("unparseable" in e["message"] for e in result["errors"])


@pytest.mark.asyncio
async def test_engine_agent_loop_respects_max_steps(sessions_root, tmp_path, docx_path):
    # Agent never says done; engine must stop at max_steps.
    steps = [
        {"action": "tool", "tool": "inspect_docx", "params": {"path": str(docx_path)}},
        {"action": "tool", "tool": "read_docx", "params": {"path": str(docx_path)}},
        {"action": "tool", "tool": "inspect_docx", "params": {"path": str(docx_path)}},
    ]
    agent = ScriptedAgent(steps)
    eng = WorkerEngine(registry, agent_factory=lambda contract: agent)
    await eng.run(make_contract(tmp_path, ["inspect_docx", "read_docx"], max_steps=2))
    await wait_done(eng)

    result = eng.get_result()
    assert result["completed_steps"] == 2
    assert len(agent.recorded) == 2  # loop cut the agent off


# ── interventions reach the agent ───────────────────────────────────────


@pytest.mark.asyncio
async def test_engine_agent_loop_routes_interventions(sessions_root, tmp_path, docx_path):
    agent = ScriptedAgent(
        [
            {"action": "tool", "tool": "inspect_docx", "params": {"path": str(docx_path)}},
            {"action": "done", "summary": "ok"},
        ]
    )
    eng = WorkerEngine(registry, agent_factory=lambda contract: agent)
    await eng.run(make_contract(tmp_path, ["inspect_docx"]))
    eng.intervene("Please also bold the title")
    await wait_done(eng)

    assert ("INTERVENTION", "Please also bold the title") in agent.recorded


# ── no factory → placeholder mode still works ───────────────────────────


@pytest.mark.asyncio
async def test_engine_without_factory_uses_placeholder(sessions_root, tmp_path):
    eng = WorkerEngine(registry)
    await eng.run(make_contract(tmp_path, ["list_directory"]))
    await wait_done(eng)

    result = eng.get_result()
    assert result["tools_used"] == ["list_directory"]
    lines = (eng._session.path / "events.jsonl").read_text(encoding="utf-8").splitlines()
    first = json.loads(lines[0])
    assert first["data"]["mode"] == "placeholder"


# ── cancel mid‑loop ─────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_engine_agent_loop_cancellable(sessions_root, tmp_path, docx_path):
    agent = ScriptedAgent(
        [
            {"action": "tool", "tool": "inspect_docx", "params": {"path": str(docx_path)}},
            {"action": "tool", "tool": "read_docx", "params": {"path": str(docx_path)}},
            {"action": "done", "summary": "late"},
        ]
    )
    eng = WorkerEngine(registry, agent_factory=lambda contract: agent)
    await eng.run(make_contract(tmp_path, ["inspect_docx", "read_docx"]))
    eng.cancel()
    await wait_done(eng)

    result = eng.get_result()
    assert result["cancelled"] is True
    assert result["success"] is False
