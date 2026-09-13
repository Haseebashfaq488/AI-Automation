import pytest
from docx import Document

from app.core.execution_engine import ExecutionEngine
from app.registry import registry
from app.workers.base.scoped_registry import ScopedToolRegistry


@pytest.fixture
def engine():
    return ExecutionEngine(registry)


@pytest.fixture
def sample_docx(tmp_path):
    """A messy .docx: markdown‑style headings, uneven content, one table."""
    doc = Document()
    doc.add_paragraph("# Introduction")
    doc.add_paragraph("Some body text here.")
    doc.add_paragraph("## Background")
    doc.add_paragraph("")
    doc.add_paragraph("More content after an empty paragraph.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    table.cell(1, 0).text = "C"
    table.cell(1, 1).text = "D"
    path = tmp_path / "sample.docx"
    doc.save(str(path))
    return path


# ── inspect_docx ────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_inspect_docx(engine, sample_docx):
    result = await engine.run("inspect_docx", {"path": str(sample_docx)})
    assert result.success
    data = result.data
    assert data["paragraph_count"] == 5
    assert data["table_count"] == 1
    assert data["heading_count"] == 0  # no real heading styles yet


@pytest.mark.asyncio
async def test_inspect_docx_missing_file(engine, tmp_path):
    result = await engine.run("inspect_docx", {"path": str(tmp_path / "nope.docx")})
    assert not result.success
    assert "not found" in result.error["message"].lower() or "Document not found" in result.error["message"]


# ── read_docx ───────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_read_docx_pagination(engine, sample_docx):
    result = await engine.run(
        "read_docx", {"path": str(sample_docx), "start": 0, "limit": 3}
    )
    assert result.success
    data = result.data
    assert data["total_paragraphs"] == 5
    assert data["returned"] == 3
    assert data["paragraphs"][0]["text"] == "# Introduction"


# ── normalize_headings ──────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_normalize_headings_dry_run(engine, sample_docx):
    result = await engine.run(
        "normalize_headings", {"path": str(sample_docx), "dry_run": True}
    )
    assert result.success
    assert result.data["changes"] == 2
    assert result.data["details"][0]["heading_level"] == 1
    assert result.data["details"][1]["heading_level"] == 2

    # dry run must NOT modify the file
    check = await engine.run("inspect_docx", {"path": str(sample_docx)})
    assert check.data["heading_count"] == 0


@pytest.mark.asyncio
async def test_normalize_headings_applies(engine, sample_docx):
    result = await engine.run("normalize_headings", {"path": str(sample_docx)})
    assert result.success
    assert result.data["changes"] == 2

    check = await engine.run("inspect_docx", {"path": str(sample_docx)})
    assert check.data["heading_count"] == 2
    headings = check.data["headings"]
    assert headings[0]["text"] == "Introduction"
    assert headings[0]["level"] == "Heading 1"
    assert headings[1]["text"] == "Background"
    assert headings[1]["level"] == "Heading 2"


# ── fix_spacing ─────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_fix_spacing_skips_empty_by_default(engine, sample_docx):
    result = await engine.run(
        "fix_spacing", {"path": str(sample_docx), "space_after": 6}
    )
    assert result.success
    # 4 non‑empty paragraphs (empty one skipped)
    assert result.data["paragraphs_affected"] == 4

    doc = Document(str(sample_docx))
    non_empty = [p for p in doc.paragraphs if p.text.strip()]
    for p in non_empty:
        assert p.paragraph_format.space_after is not None


@pytest.mark.asyncio
async def test_fix_spacing_include_empty(engine, sample_docx):
    result = await engine.run(
        "fix_spacing", {"path": str(sample_docx), "include_empty": True}
    )
    assert result.success
    assert result.data["paragraphs_affected"] == 5


# ── format_tables ───────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_format_tables(engine, sample_docx):
    result = await engine.run(
        "format_tables", {"path": str(sample_docx), "style": "Table Grid"}
    )
    assert result.success
    assert result.data["tables_affected"] == 1

    doc = Document(str(sample_docx))
    assert doc.tables[0].style.name == "Table Grid"


@pytest.mark.asyncio
async def test_format_tables_unknown_style(engine, sample_docx):
    result = await engine.run(
        "format_tables", {"path": str(sample_docx), "style": "No Such Style"}
    )
    assert result.success  # tool handles it gracefully
    assert "error" in result.data
    assert "Unknown table style" in result.data["error"]


# ── backup_docx ─────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_backup_docx(engine, sample_docx):
    result = await engine.run("backup_docx", {"path": str(sample_docx)})
    assert result.success
    backup = result.data["backup"]
    assert backup.endswith(".backup.docx")
    import os
    assert os.path.isfile(backup)

    # Backup is identical content
    original = Document(str(sample_docx))
    copied = Document(backup)
    assert len(original.paragraphs) == len(copied.paragraphs)


# ── fs_scope enforcement ────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_fs_scope_blocks_outside_path(engine, sample_docx, tmp_path):
    other = tmp_path / "other"
    other.mkdir()
    result = await engine.run(
        "inspect_docx",
        {"path": str(sample_docx), "fs_scope": str(other)},
    )
    assert not result.success
    assert "outside" in result.error["message"]


@pytest.mark.asyncio
async def test_fs_scope_allows_inside_path(engine, sample_docx, tmp_path):
    result = await engine.run(
        "inspect_docx",
        {"path": str(sample_docx), "fs_scope": str(tmp_path)},
    )
    assert result.success


# ── create_docx ─────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_create_docx_basic(engine, tmp_path):
    new_doc = tmp_path / "new_doc.docx"
    result = await engine.run(
        "create_docx",
        {
            "path": str(new_doc),
            "title": "Project Alpha Report",
            "initial_text": "This document contains project metrics.",
        },
    )
    assert result.success
    assert result.data["created"] is True
    assert result.data["paragraph_count"] >= 2

    # Verify directly via python-docx
    doc = Document(str(new_doc))
    assert doc.paragraphs[0].text == "Project Alpha Report"
    assert doc.paragraphs[1].text == "This document contains project metrics."


@pytest.mark.asyncio
async def test_create_docx_overwrite_protection(engine, tmp_path):
    existing = tmp_path / "exists.docx"
    doc = Document()
    doc.add_paragraph("Original content")
    doc.save(str(existing))

    # Without overwrite=True, should fail with FileExistsError
    res1 = await engine.run("create_docx", {"path": str(existing), "overwrite": False})
    assert not res1.success
    assert "already exists" in res1.error["message"].lower()

    # With overwrite=True, should succeed
    res2 = await engine.run(
        "create_docx",
        {"path": str(existing), "title": "Overwritten", "overwrite": True},
    )
    assert res2.success
    updated = Document(str(existing))
    assert updated.paragraphs[0].text == "Overwritten"


@pytest.mark.asyncio
async def test_create_docx_fs_scope_protection(engine, tmp_path):
    scope = tmp_path / "allowed_folder"
    scope.mkdir()
    outside = tmp_path / "forbidden.docx"

    result = await engine.run(
        "create_docx",
        {"path": str(outside), "fs_scope": str(scope)},
    )
    assert not result.success
    assert "outside" in result.error["message"]


# ── add_heading ──────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_add_heading(engine, sample_docx):
    result = await engine.run(
        "add_heading",
        {"path": str(sample_docx), "text": "New Findings", "level": 2},
    )
    assert result.success
    assert result.data["added_heading"] == "New Findings"
    assert result.data["level"] == 2

    doc = Document(str(sample_docx))
    last_para = doc.paragraphs[-1]
    assert last_para.text == "New Findings"
    assert last_para.style.name.startswith("Heading")


# ── add_paragraph ────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_add_paragraph_formatted(engine, sample_docx):
    result = await engine.run(
        "add_paragraph",
        {
            "path": str(sample_docx),
            "text": "Critical warning details",
            "bold": True,
            "italic": True,
        },
    )
    assert result.success
    assert result.data["bold"] is True
    assert result.data["italic"] is True

    doc = Document(str(sample_docx))
    last_para = doc.paragraphs[-1]
    assert last_para.text == "Critical warning details"
    run = last_para.runs[0]
    assert run.bold is True
    assert run.italic is True


# ── add_table ────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_add_table(engine, sample_docx):
    headers = ["Task", "Owner", "Status"]
    rows = [
        ["Setup backend", "Alice", "Done"],
        ["Design frontend", "Bob", "In Progress"],
    ]
    result = await engine.run(
        "add_table",
        {
            "path": str(sample_docx),
            "headers": headers,
            "rows": rows,
        },
    )
    assert result.success
    assert result.data["rows_count"] == 2
    assert result.data["cols_count"] == 3

    doc = Document(str(sample_docx))
    assert len(doc.tables) == 2  # sample_docx had 1 table + new table
    new_table = doc.tables[-1]
    assert len(new_table.rows) == 3
    assert new_table.cell(0, 0).text == "Task"
    assert new_table.cell(1, 0).text == "Setup backend"
    assert new_table.cell(1, 2).text == "Done"


@pytest.mark.asyncio
async def test_add_table_inferred_headers(engine, sample_docx):
    # Missing 'headers' parameter — inferred from row 0
    rows = [
        ["Metric", "Value"],
        ["CPU", "45%"],
        ["Memory", "60%"],
    ]
    result = await engine.run(
        "add_table",
        {
            "path": str(sample_docx),
            "rows": rows,
        },
    )
    assert result.success
    assert result.data["headers"] == ["Metric", "Value"]
    assert result.data["rows_count"] == 2


@pytest.mark.asyncio
async def test_add_table_dict_rows(engine, sample_docx):
    # Rows as list of dictionaries
    records = [
        {"Service": "Auth", "Status": "Healthy"},
        {"Service": "Database", "Status": "Warning"},
    ]
    result = await engine.run(
        "add_table",
        {
            "path": str(sample_docx),
            "rows": records,
        },
    )
    assert result.success
    assert "Service" in result.data["headers"]
    assert "Status" in result.data["headers"]

    doc = Document(str(sample_docx))
    new_table = doc.tables[-1]
    assert new_table.cell(1, 0).text == "Auth"
    assert new_table.cell(2, 1).text == "Warning"


@pytest.mark.asyncio
async def test_add_table_string_list(engine, sample_docx):
    # Rows as 1D list of text paragraphs
    items = ["Summary point 1", "Summary point 2"]
    result = await engine.run(
        "add_table",
        {
            "path": str(sample_docx),
            "headers": ["Findings"],
            "rows": items,
        },
    )
    assert result.success
    assert result.data["rows_count"] == 2
    assert result.data["cols_count"] == 1

    doc = Document(str(sample_docx))
    new_table = doc.tables[-1]
    assert new_table.cell(1, 0).text == "Summary point 1"
    assert new_table.cell(2, 0).text == "Summary point 2"


@pytest.mark.asyncio
async def test_add_table_string_headers_and_aliases(engine, sample_docx):
    # Using 'columns' string alias and 'data' alias
    result = await engine.run(
        "add_table",
        {
            "path": str(sample_docx),
            "columns": "ID, Name, Role",
            "data": [["1", "Alice", "Admin"]],
        },
    )
    assert result.success
    assert result.data["headers"] == ["ID", "Name", "Role"]
    assert result.data["cols_count"] == 3



# ── scoped worker view only exposes word tools given to it ──────────────


def test_word_tools_visible_in_scoped_registry():
    all_word_tools = {
        "create_docx", "add_heading", "add_paragraph", "add_table",
        "inspect_docx", "read_docx", "normalize_headings",
        "fix_spacing", "format_tables", "backup_docx",
    }
    scoped = ScopedToolRegistry(registry, all_word_tools)
    tools = scoped.list_tools()
    assert set(tools.keys()) == all_word_tools

