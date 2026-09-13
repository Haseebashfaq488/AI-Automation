from pathlib import Path
from typing import Any, Dict, Optional

from docx import Document

from app.core.tool import BaseTool, RiskLevel
from app.workers.document_worker.word_worker.helpers.document import save_docx
from app.workers.document_worker.word_worker.helpers.paths import (
    ensure_docx_extension,
    resolve_doc_path,
)


class CreateDocxTool(BaseTool):
    """Create a new Word document (.docx) with optional title and content.

    Parameters
    ----------
    path: str
        Absolute path where the .docx file will be created.
    title: str (optional)
        Main document title to insert at the top (Level 0 / Title style).
    initial_text: str (optional)
        Initial paragraph text to insert.
    overwrite: bool (optional)
        Whether to overwrite if file already exists (default: False).
    fs_scope: str (optional)
        Worker scope the path must reside within.
    """

    name = "create_docx"
    description = "Create a new Word document (.docx) with an optional title and initial content."
    risk = RiskLevel.MEDIUM
    input_schema = {
        "type": "object",
        "properties": {
            "path": {"type": "string"},
            "title": {"type": "string"},
            "initial_text": {"type": "string"},
            "overwrite": {"type": "boolean", "default": False},
            "fs_scope": {"type": "string"},
        },
        "required": ["path"],
    }

    async def execute(self, params: Dict[str, Any]) -> Dict[str, Any]:
        path = resolve_doc_path(params["path"], params.get("fs_scope"))
        ensure_docx_extension(path)

        overwrite = params.get("overwrite", False)
        if path.exists() and not overwrite:
            raise FileExistsError(f"File already exists: {path}. Set overwrite=True to replace it.")

        path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()
        title = params.get("title")
        if title:
            doc.add_heading(str(title).strip(), level=0)

        initial_text = params.get("initial_text")
        if initial_text:
            doc.add_paragraph(str(initial_text))

        save_docx(doc, path)

        return {
            "path": str(path),
            "created": True,
            "title": title or None,
            "paragraph_count": len(doc.paragraphs),
            "size_bytes": path.stat().st_size,
        }
