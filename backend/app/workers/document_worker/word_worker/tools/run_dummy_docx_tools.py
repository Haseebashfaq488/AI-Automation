import json
import sys
import asyncio
from pathlib import Path

# Add backend root to sys.path so that 'app' package can be imported
backend_root = Path(__file__).resolve().parents[5]
if str(backend_root) not in sys.path:
    sys.path.insert(0, str(backend_root))

# Ensure python-docx is available
try:
    from docx import Document
except ImportError:
    raise RuntimeError('python-docx library is required but not installed.')

# Import tools using absolute package path
from app.workers.document_worker.word_worker.tools.read_docx import ReadDocxTool
from app.workers.document_worker.word_worker.tools.inspect_docx import InspectDocxTool

async def main():
    # Create dummy docx file in the same directory
    dummy_path = Path(__file__).parent / 'dummy.docx'
    if dummy_path.exists():
        dummy_path.unlink()

    doc = Document()
    doc.add_heading('Test Document', level=1)
    doc.add_paragraph('This is a paragraph in the dummy document.')
    doc.add_paragraph('Another paragraph with default style.')
    doc.save(dummy_path)

    # Parameters for ReadDocxTool
    read_params = {
        "path": str(dummy_path.resolve()),
        "fs_scope": None,
        "start": 0,
        "limit": 200,
    }

    read_tool = ReadDocxTool()
    read_result = await read_tool.execute(read_params)

    # Parameters for InspectDocxTool
    inspect_params = {"path": str(dummy_path.resolve()), "fs_scope": None}
    inspect_tool = InspectDocxTool()
    inspect_result = await inspect_tool.execute(inspect_params)

    output = {
        "read_docx": read_result,
        "inspect_docx": inspect_result,
    }
    print(json.dumps(output, indent=2))

if __name__ == "__main__":
    asyncio.run(main())
