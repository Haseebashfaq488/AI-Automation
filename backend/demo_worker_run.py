"""demo_worker_run.py -- full worker run with live terminal output.

Creates a messy report.docx and forks a real worker, printing every
decision and step to the terminal as it happens. At the end it verifies
the results on disk.

Usage (from the backend directory):

    python demo_worker_run.py          # scripted agent: deterministic, offline
    python demo_worker_run.py --live   # real LLM agent (Groq key in .env)
    python demo_worker_run.py --live --model llama-3.1-8b-instant

The scripted run is instant; the live run takes a few minutes because each
step is one LLM call. Both are safe to run any number of times -- the demo
workspace and worker session are recreated each run.
"""

import asyncio
import json
import os
import sys
import time
from pathlib import Path

# Make the backend package importable regardless of CWD
BACKEND_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(BACKEND_DIR))


def _parse_args():
    """Extract --live and --model <id>; the model becomes WORKER_MODEL."""
    args = sys.argv[1:]
    live = "--live" in args
    model = None
    if "--model" in args:
        i = args.index("--model")
        if i + 1 < len(args):
            model = args[i + 1]
            os.environ["WORKER_MODEL"] = model  # read by get_worker_model()
    return live, model


LIVE, CLI_MODEL = _parse_args()

from docx import Document  # noqa: E402

from app.registry import registry  # noqa: E402
from app.workers.base import bus as bus_mod  # noqa: E402
from app.workers.base.contract import TaskContract  # noqa: E402
from app.workers.base.engine import WorkerEngine  # noqa: E402

WORK_DIR = BACKEND_DIR / "demo_workspace"
DOCX_PATH = WORK_DIR / "report.docx"

# The exact sequence the scripted "agent" will decide (offline mode).
SCRIPT = [
    ("backup_docx", {"path": str(DOCX_PATH)}),
    ("inspect_docx", {"path": str(DOCX_PATH)}),
    ("normalize_headings", {"path": str(DOCX_PATH)}),
    ("fix_spacing", {"path": str(DOCX_PATH), "space_after": 6}),
    ("format_tables", {"path": str(DOCX_PATH), "style": "Table Grid"}),
    ("inspect_docx", {"path": str(DOCX_PATH)}),  # verify
]


class ScriptedAgent:
    """Pretends to be the LLM: replays SCRIPT, reports everything it does."""

    def __init__(self):
        self.decisions = [dict(tool=t, params=p) for t, p in SCRIPT]

    # What the engine checks to pick the agent loop over the placeholder loop
    available = True

    async def decide_next_step(self):
        if not self.decisions:
            return {
                "action": "done",
                "summary": "Backed up, converted markdown headings to real Word "
                           "heading styles, fixed spacing, applied Table Grid, "
                           "and verified by re-inspecting the document.",
            }
        nxt = self.decisions.pop(0)
        print(f"[agent ] decision -> call {nxt['tool']} {nxt['params']}")
        return {"action": "tool", "tool": nxt["tool"], "params": nxt["params"]}

    def record_step(self, tool, success, output=None, error=None):
        if success:
            print(f"[agent ] observed: {tool} ok")
        else:
            print(f"[agent ] observed: {tool} FAILED ({error})")

    def inject_intervention(self, message):
        print(f"[agent ] parent says: {message}")


class LiveAgent:
    """Wraps the real DocumentWorkerAgent, printing its decisions/observations."""

    def __init__(self, inner):
        self._inner = inner
        self.available = inner.available

    async def decide_next_step(self):
        started = time.monotonic()
        decision = await self._inner.decide_next_step()
        elapsed = time.monotonic() - started
        action = decision.get("action")
        if action == "tool":
            print(f"[agent ] LLM decision -> call {decision['tool']} {decision['params']}"
                  f"  (took {elapsed:.1f}s)")
        elif action == "done":
            print(f"[agent ] LLM says DONE: {decision.get('summary', '')}  (took {elapsed:.1f}s)")
        else:
            print(f"[agent ] LLM decision FAILED: {decision.get('reason', '?')}  (took {elapsed:.1f}s)")
        return decision

    def record_step(self, tool, success, output=None, error=None):
        if success:
            print(f"[agent ] observed: {tool} ok")
        else:
            print(f"[agent ] observed: {tool} FAILED ({error})")
        self._inner.record_step(tool, success, output, error)

    def inject_intervention(self, message):
        print(f"[agent ] parent says: {message}")
        self._inner.inject_intervention(message)


def make_agent_factory():
    """Scripted (offline) factory by default; real LLM agent with --live."""
    if not LIVE:
        agent = ScriptedAgent()
        return lambda contract: agent

    from app.workers.document_worker.agent.worker_agent import DocumentWorkerAgent
    from app.workers.document_worker.agent.config import get_worker_model

    probe = DocumentWorkerAgent("probe", ["inspect_docx"], str(WORK_DIR))
    if not probe.available:
        print("ERROR: no worker LLM key found. Set GROQ_API_KEY in backend/.env for --live,")
        print("       or run without --live for the scripted demo.")
        sys.exit(1)
    print(f"[setup ] live agent: LLM model {get_worker_model()}" + (" (from --model)" if CLI_MODEL else " (default)"))

    def factory(contract):
        return LiveAgent(
            DocumentWorkerAgent(contract.objective, contract.allowed_tools, contract.fs_scope)
        )

    return factory


def make_messy_docx() -> None:
    """A document with markdown-style headings, no spacing, unstyled table."""
    WORK_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_paragraph("# Quarterly Report")
    doc.add_paragraph("Intro paragraph with default spacing.")
    doc.add_paragraph("## Sales Figures")
    doc.add_paragraph("Some body text.")
    doc.add_paragraph("## Outlook")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "North"
    table.cell(1, 1).text = "1200"
    doc.save(str(DOCX_PATH))


def subscribe_terminal(session_id: str) -> None:
    """Print every worker event live as the engine emits it."""

    def _cb(event_type: str, data: dict) -> None:
        if event_type == "WORK_STARTED":
            print(f"[worker] WORK STARTED -- objective: {data.get('objective')}")
            print(f"[worker] tools allowed: {', '.join(data.get('tools', []))}")
        elif event_type == "STEP_STARTED":
            print(f"[worker]   step {data.get('index', '?') + 1} started: {data.get('step')}")
        elif event_type == "STEP_COMPLETED":
            out = data.get("output") or {}
            brief = {k: v for k, v in out.items() if isinstance(v, (int, str, bool))}
            print(f"[worker]   step done:   {data.get('step')} -> {brief}")
        elif event_type == "STEP_FAILED":
            print(f"[worker]   step FAILED: {data.get('step')} -- {data.get('error')}")
        elif event_type == "WORK_COMPLETED":
            print(f"[worker] WORK COMPLETED -- success={data.get('success')}, "
                  f"steps={data.get('completed_steps')}, cancelled={data.get('cancelled')}")

    bus_mod.on(session_id, _cb)


async def run() -> None:
    mode = "LIVE LLM agent" if LIVE else "scripted agent (offline)"
    print(f"=== Jarvis worker demo: {mode} ===")
    make_messy_docx()
    print(f"[setup ] messy document created: {DOCX_PATH}")

    contract = TaskContract(
        objective="Clean up report.docx: convert the markdown-style # headings "
                  "to real Word heading styles, set consistent paragraph spacing, "
                  "apply a proper table style, and verify by re-inspecting.",
        requirements=[],
        constraints=[],
        success_criteria=["Headings converted", "Backup created before edits"],
        fs_scope=str(WORK_DIR),
        allowed_tools=["backup_docx", "inspect_docx", "read_docx",
                       "normalize_headings", "fix_spacing", "format_tables"],
        max_steps=10,
    )

    engine = WorkerEngine(registry, agent_factory=make_agent_factory())
    subscribe_terminal(engine.session_id())

    print(f"[setup ] worker forked: {engine.session_id()}")
    print(f"[setup ] fs_scope (worker may only touch this folder): {WORK_DIR}")
    print("[setup ] starting loop... (Ctrl+C to abort)\n")

    await engine.run(contract)

    # Poll until the background loop terminates
    last_status = None
    while True:
        state = engine.get_state()
        status_line = (state["status"], state["progress_percent"], state["current_step"])
        if status_line != last_status:
            print(f"[poll  ] {state['status']} {state['progress_percent']}%"
                  + (f" -- step: {state['current_step']}" if state["current_step"] else ""))
            last_status = status_line
        if state["status"] in ("completed", "cancelled"):
            break
        await asyncio.sleep(0.5)

    print("\n=== VERIFYING RESULTS ON DISK ===")
    doc = Document(str(DOCX_PATH))
    headings = [(p.style.name, p.text) for p in doc.paragraphs
                if p.style.name.startswith("Heading")]
    print(f"report.docx headings now: {headings}")
    print(f"report.docx table style:  {doc.tables[0].style.name}")

    backup = WORK_DIR / "report.backup.docx"
    print(f"backup exists: {backup.exists()} ({backup.stat().st_size} bytes)" if backup.exists()
          else "backup exists: NO -- MISSING!")

    result = engine.get_result()
    print(f"result.json: success={result.get('success')} "
          f"steps={result.get('completed_steps')} errors={result.get('errors')}")
    print(f"session on disk: {engine._session.path}")
    print(f"event log:       {engine._session.path / 'events.jsonl'}")
    print("\nDemo finished. Delete demo_workspace\\ and the session folder to reset.")


if __name__ == "__main__":
    try:
        asyncio.run(run())
    except KeyboardInterrupt:
        print("\nInterrupted.")
