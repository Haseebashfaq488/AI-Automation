import json
import sys
from pathlib import Path

# Add project root to sys.path so that 'app' package can be imported
project_root = Path(__file__).resolve().parent
if str(project_root) not in sys.path:
    sys.path.insert(0, str(project_root))

# Ensure python-docx is installed
try:
    from docx import Document
except ImportError:
    raise RuntimeError('python-docx library is required but not installed.')

# Import tools using absolute package path
from app.workers.document_worker.word_worker.tools.read_docx import ReadDocxTool
from app.workers.document_worker.word_worker.tools.inspect_docx import InspectDocxTool

# Create dummy docx file
dummy_path = project_root / 'dummy.docx'
if dummy_path.exists():
    dummy_path.unlink()

doc = Document()
doc.add_heading('Test Document', level=1)
doc.add_paragraph('This is a paragraph in the dummy document.')
doc.add_paragraph('Another paragraph with default style.')

doc.save(dummy_path)

# Prepare parameters for ReadDocxTool
read_params = {
    "path": str(dummy_path.resolve()),
    "fs_scope": None,
    "start": 0,
    "limit": 200,
}

read_tool = ReadDocxTool()
read_result = read_tool.execute(read_params)

# Prepare parameters for InspectDocxTool
inspect_params = {"path": str(dummy_path.resolve()), "fs_scope": None}
inspect_tool = InspectDocxTool()
inspect_result = inspect_tool.execute(inspect_params)

# Output results as JSON
output = {
    "read_docx": read_result,
    "inspect_docx": inspect_result,
}
print(json.dumps(output, indent=2))
